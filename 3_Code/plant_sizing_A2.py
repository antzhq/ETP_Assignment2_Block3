"""
Hybrid CSP + PVT + wind + biomass plant for Santiago (Assignment 2, Questions 6 to 10).

This script implements the mathematical model in "../2_Math_Model/Q6-Q10_Mathematical_Model.pdf" one-to-one.
Every variable carries the same name as its symbol in that document, e.g.
    Use_{wind,hr,day}        ->  use_wind[day][hr]
    Count_{turbines}         ->  count_turbines

It extends the Assignment 1 model (plant_sizing.py). The CSP and PVT physics are unchanged.

What it does, in order:
    1. Reads the hourly data: irradiance, ambient temperature, demand (Assignment 1 file)
       and wind speed (Assignment 2 file, km/h converted to m/s).
    2. Calculates the derived parameters: energy per heliostat, per PVT panel and per wind turbine
       in every hour, and the ethanol needed per GJ of electricity.
    3. Solves the model as a mixed-integer linear program (scipy / HiGHS):
         minimise   unmet demand (heavily weighted) + surplus energy
         subject to hourly electricity balance, CSP produces >= 60% of ALL electricity
                    produced each day (surplus included),
                    wind supplies at least as much electricity as ethanol on each day,
                    PVT heat as in Assignment 1, site area.
    4. Prints the answers to Q6, Q7 and Q8 and saves four figures to ../4_Visualizations:
         Q9_Hourly_Supply_vs_Demand.png        hourly dispatch vs demand, both days, electricity and heat
         Q9_Comparison_with_Assignment_1.png   daily energy balance, Assignment 1 plant vs this plant
         Q10_Wind_vs_Ethanol.png               wind vs ethanol produced, same plant, 75-95% of Betz
         Q10_Wind_Efficiency_Sweep.png         total electricity produced by source, 75-95% of the Betz limit
    5. Exports all numbers to ../5_Data/Results_Hybrid_Plant_A2.xlsx.

How to run (paths are resolved relative to this file):
    pip install numpy scipy matplotlib openpyxl
    python plant_sizing_A2.py
"""

import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.patches import Patch
from scipy.optimize import Bounds, LinearConstraint, milp

# =============================================================================
# 1. SETS
# =============================================================================
days = ["June", "December"]  # Days
hrs = list(range(24))         # Hrs = {0, 1, ..., 23}

# =============================================================================
# 2. PARAMETERS
# =============================================================================
# --- CSP (Assignment 1, unchanged) ---
eff_reflectivity = 0.95
eff_cosine = 0.86
eff_shading_blocking = 0.89
eff_atmospheric = 0.97
eff_spillage = 0.92
eff_receiver = 0.80
eff_rankine = 0.42            # Rankine cycle thermal efficiency (solar AND ethanol firing)
area_per_heliostat = 220.0    # m2

# --- PVT (Assignment 1, unchanged) ---
eff_PVT_elec = 0.16
frac_absorbed = 0.86
coef_heat_loss = 12.0         # W/(m2 K)
factor_heat_removal = 0.65
T_water_in = 35.0             # degC
area_per_PVT_panel = 4.0      # m2

# --- Wind (Assignment 2) ---
betz_limit = 16 / 27                  # maximum power coefficient of any turbine
frac_of_betz = 0.85                   # turbine efficiency as a fraction of the Betz limit (Q10 varies 0.75-0.95)
radius_rotor = 36.5                   # m
density_air = 1.225                   # kg/m3 at sea level (Lecture 6; not given in the brief -> assumption)
area_rotor = math.pi * radius_rotor ** 2  # swept area, m2
kmh_per_ms = 3.6                      # 1 m/s = 3.6 km/h

# --- Biomass: ethanol (Assignment 2, Q2-Q5) ---
HHV_ethanol = 29.7                    # MJ/kg
H_evap_water = 2.5                    # MJ/kg of water
molar_mass_ethanol = 46.07            # kg/kmol (same values as Q5)
molar_mass_water = 18.015             # kg/kmol
kmol_water_per_kmol_ethanol = 3       # from the balanced reaction (Q2)
mass_water_per_kg_ethanol = kmol_water_per_kmol_ethanol * molar_mass_water / molar_mass_ethanol
LHV_ethanol = HHV_ethanol - mass_water_per_kg_ethanol * H_evap_water  # MJ/kg, heat actually delivered (Q5)

# --- Plant and rules ---
area_max = 450_000.0                  # m2 of collectors (mirrors + PVT); turbines stand outside this site
share_CSP_min = 0.60                  # CSP produces >= 60% of all electricity produced that day (surplus included)
weight_unmet = 1000.0                 # penalty per GJ of unmet demand (>> 1 per GJ of surplus)
weight_merit = 1e-4                   # tiny tie-breaker for the dispatch order wind > CSP > PVT (see model)
weight_biomass = 1e-3                 # tiny tie-breaker: burn ethanol only when nothing else can serve demand
GJ_per_W_hr = 3600 / 1e9              # energy of 1 W for 1 hour, GJ
MJ_per_GJ = 1000.0

# --- Assignment 1 design (for the comparison in Q8 and Q9) ---
count_heliostats_A1 = 215
count_PVT_panels_A1 = 9_559


# =============================================================================
# 2b. DATA
# =============================================================================
def read_data(path_A1, path_wind):
    """Hourly data, indexed variable[day][hr]. Wind speed is converted from km/h to m/s."""
    workbook = openpyxl.load_workbook(path_A1, data_only=True)
    sheet_supply, sheet_demand = workbook["Supply"], workbook["Demand"]
    columns = {
        "June":     {"irradiance": "B", "T_ambient": "C", "demand_heat": "B", "demand_elec": "D", "wind": "C"},
        "December": {"irradiance": "E", "T_ambient": "F", "demand_heat": "K", "demand_elec": "M", "wind": "E"},
    }
    sheet_wind = openpyxl.load_workbook(path_wind, data_only=True)["wind"]
    first_row_A1, first_row_wind = 4, 6

    irradiance, T_ambient, demand_heat, demand_elec, speed_wind = {}, {}, {}, {}, {}
    for day in days:
        col = columns[day]
        rows = [first_row_A1 + hr for hr in hrs]
        irradiance[day] = np.array([float(sheet_supply[f"{col['irradiance']}{r}"].value) for r in rows])
        T_ambient[day] = np.array([float(sheet_supply[f"{col['T_ambient']}{r}"].value) for r in rows])
        demand_heat[day] = np.array([float(sheet_demand[f"{col['demand_heat']}{r}"].value) for r in rows])
        demand_elec[day] = np.array([float(sheet_demand[f"{col['demand_elec']}{r}"].value) for r in rows])
        speed_wind[day] = np.array([float(sheet_wind[f"{col['wind']}{first_row_wind + hr}"].value)
                                    for hr in hrs]) / kmh_per_ms
    return irradiance, T_ambient, demand_heat, demand_elec, speed_wind


# =============================================================================
# 3. DERIVED PARAMETERS (energy per unit in every hour, GJ)
# =============================================================================
def derived(data, frac_betz=frac_of_betz):
    """
    Energy one unit of each technology can deliver in each hour, in GJ:
        energy_per_heliostat[day][hr]      electricity from one heliostat (Assignment 1 chain)
        energy_per_PVT_elec[day][hr]       electricity from one PVT panel
        energy_per_PVT_heat[day][hr]       useful heat from one PVT panel (Assignment 1 formula)
        energy_per_turbine[day][hr]        electricity from one wind turbine:
                                               P = 1/2 * rho * A * v^3 * Cp,   Cp = frac_betz * 16/27
    """
    irradiance, T_ambient, _, _, speed_wind = data
    eff_optical = eff_reflectivity * eff_cosine * eff_shading_blocking * eff_atmospheric * eff_spillage
    eff_CSP = eff_optical * eff_receiver * eff_rankine
    coef_power = frac_betz * betz_limit
    out = {"eff_CSP": eff_CSP, "coef_power": coef_power,
           "energy_per_heliostat": {}, "energy_per_PVT_elec": {}, "energy_per_PVT_heat": {},
           "energy_per_turbine": {}, "sun_hrs": {}, "wind_hrs": {}}
    for day in days:
        out["energy_per_heliostat"][day] = GJ_per_W_hr * area_per_heliostat * eff_CSP * irradiance[day]
        out["energy_per_PVT_elec"][day] = GJ_per_W_hr * area_per_PVT_panel * eff_PVT_elec * irradiance[day]
        yield_heat = np.maximum(0.0, factor_heat_removal * (
            frac_absorbed * irradiance[day] - coef_heat_loss * (T_water_in - T_ambient[day])))
        out["energy_per_PVT_heat"][day] = GJ_per_W_hr * area_per_PVT_panel * yield_heat
        power_turbine = 0.5 * density_air * area_rotor * speed_wind[day] ** 3 * coef_power  # W
        out["energy_per_turbine"][day] = GJ_per_W_hr * power_turbine
        out["sun_hrs"][day] = irradiance[day] > 0
        out["wind_hrs"][day] = speed_wind[day] > 0
    return out


def ethanol_kg(use_biomass_GJ):
    """kg of ethanol to make the given electricity [GJ]: E_elec = m * LHV * eff_rankine."""
    return use_biomass_GJ * MJ_per_GJ / (LHV_ethanol * eff_rankine)


# =============================================================================
# 5. MODEL (mixed-integer linear program)
# =============================================================================
# Variable layout in the solver vector x:
#   x[0] Count_heliostats, x[1] Count_PVT_panels, x[2] Count_turbines,
#   then for each day and hour: Use_CSP, Use_PVT, Use_wind, Use_biomass, Unmet
SOURCES = ["CSP", "PVT", "wind", "biomass", "unmet"]


def index(day_i, hr, source):
    return 3 + day_i * 24 * len(SOURCES) + hr * len(SOURCES) + SOURCES.index(source)


def solve(data, par, fixed_counts=None, fixed_turbines=None, wind_rule=True, share_rule=True):
    """
    Builds and solves the model. Returns counts and the hourly dispatch.
    fixed_counts   = (heliostats, PVT panels, turbines) fixes the whole design and only dispatches (Q10a).
    fixed_turbines = number of turbines fixed, the rest optimised (trade-off curve).
    wind_rule      = apply (C5) wind >= ethanol on each day.
    share_rule     = apply (C4) CSP >= 60% of all electricity produced.
    """
    _, _, demand_heat, demand_elec, _ = data
    n = 3 + 2 * 24 * len(SOURCES)
    rows, lower, upper = [], [], []

    def add(row, lo, hi):
        rows.append(row); lower.append(lo); upper.append(hi)

    # ---- objective: weight_unmet * Unmet + Surplus_elec + Surplus_heat (constants dropped) ----
    c = np.zeros(n)
    for i, day in enumerate(days):
        c[0] += par["energy_per_heliostat"][day].sum()          # available CSP electricity
        c[1] += par["energy_per_PVT_elec"][day].sum()           # available PVT electricity
        c[1] += par["energy_per_PVT_heat"][day][par["sun_hrs"][day]].sum()  # PVT heat (heat surplus)
        c[2] += par["energy_per_turbine"][day].sum()            # available wind electricity
        for hr in hrs:
            # used energy is not surplus; tiny extra weights set the dispatch order when several
            # renewables could serve the same hour: wind first (preference to wind), then CSP
            # (its 60% is a hard constraint anyway), PVT last (it is built for heat)
            c[index(i, hr, "wind")] = -1 - 3 * weight_merit
            c[index(i, hr, "CSP")] = -1 - 2 * weight_merit
            c[index(i, hr, "PVT")] = -1 - 1 * weight_merit
            c[index(i, hr, "unmet")] = weight_unmet
            c[index(i, hr, "biomass")] = weight_biomass

    for i, day in enumerate(days):
        sun = par["sun_hrs"][day]
        # (C1) heat demand over the sunlit hours (Assignment 1)
        row = np.zeros(n); row[1] = par["energy_per_PVT_heat"][day][sun].sum()
        add(row, demand_heat[day][sun].sum(), np.inf)
        for hr in hrs:
            # (C2) hourly electricity balance
            row = np.zeros(n)
            for s in SOURCES:
                row[index(i, hr, s)] = 1
            add(row, demand_elec[day][hr], demand_elec[day][hr])
            # (C3) a source cannot deliver more than it generates
            for s, k, key in [("CSP", 0, "energy_per_heliostat"), ("PVT", 1, "energy_per_PVT_elec"),
                              ("wind", 2, "energy_per_turbine")]:
                row = np.zeros(n); row[index(i, hr, s)] = 1; row[k] = -par[key][day][hr]
                add(row, -np.inf, 0)
        # (C4) CSP produces >= 60% of ALL electricity produced that day (surplus included):
        #      Gen_CSP >= Share * (Gen_CSP + Gen_PVT + Gen_wind + Use_biomass)
        #  <=> (1 - Share) * Gen_CSP - Share * (Gen_PVT + Gen_wind + Use_biomass) >= 0
        if share_rule:
            row = np.zeros(n)
            row[0] = (1 - share_CSP_min) * par["energy_per_heliostat"][day].sum()
            row[1] = -share_CSP_min * par["energy_per_PVT_elec"][day].sum()
            row[2] = -share_CSP_min * par["energy_per_turbine"][day].sum()
            for hr in hrs:
                row[index(i, hr, "biomass")] = -share_CSP_min
            add(row, 0, np.inf)
        # (C5) preference to wind: over the day, wind delivers at least as much electricity as ethanol
        if wind_rule:
            row = np.zeros(n)
            for hr in hrs:
                row[index(i, hr, "wind")] = 1
                row[index(i, hr, "biomass")] = -1
            add(row, 0, np.inf)
    # (C6) site area
    row = np.zeros(n); row[0] = area_per_heliostat; row[1] = area_per_PVT_panel
    add(row, -np.inf, area_max)

    lb, ub = np.zeros(n), np.full(n, np.inf)
    if fixed_turbines is not None:
        lb[2] = ub[2] = fixed_turbines
    if fixed_counts is not None:
        lb[:3] = ub[:3] = fixed_counts
    integrality = np.zeros(n); integrality[:3] = 1  # (C7) whole units

    res = milp(c, constraints=LinearConstraint(np.array(rows), lower, upper),
               integrality=integrality, bounds=Bounds(lb, ub))
    if not res.success:
        raise RuntimeError(res.message)
    x = res.x
    counts = [int(round(v)) for v in x[:3]]
    use = {s: {day: np.array([x[index(i, hr, s)] for hr in hrs]) for i, day in enumerate(days)} for s in SOURCES}
    return counts, use


def summarise(counts, use, data, par):
    """Everything needed for printing, plotting and Excel, for one design."""
    _, _, demand_heat, demand_elec, _ = data
    count_heliostats, count_PVT_panels, count_turbines = counts
    r = {"counts": counts, "use": use, "par": par,
         "avail_CSP": {}, "avail_PVT": {}, "avail_wind": {}, "supply_heat": {},
         "surplus_elec": {}, "ethanol_kg": {}}
    for day in days:
        r["avail_CSP"][day] = count_heliostats * par["energy_per_heliostat"][day]
        r["avail_PVT"][day] = count_PVT_panels * par["energy_per_PVT_elec"][day]
        r["avail_wind"][day] = count_turbines * par["energy_per_turbine"][day]
        r["supply_heat"][day] = count_PVT_panels * par["energy_per_PVT_heat"][day]
        r["surplus_elec"][day] = (r["avail_CSP"][day] + r["avail_PVT"][day] + r["avail_wind"][day]
                                  - use["CSP"][day] - use["PVT"][day] - use["wind"][day])
        r["ethanol_kg"][day] = ethanol_kg(use["biomass"][day])
    return r


# =============================================================================
# 6. ASSIGNMENT 1 PLANT ON THE SAME DATA (for the comparison)
# =============================================================================
def assignment1_balance(data, par):
    """Hourly supply of the Assignment 1 plant (215 heliostats, 9,559 PVT panels, no wind or ethanol)."""
    _, _, _, demand_elec, _ = data
    out = {}
    for day in days:
        csp = count_heliostats_A1 * par["energy_per_heliostat"][day]
        pvt = count_PVT_panels_A1 * par["energy_per_PVT_elec"][day]
        used_csp = np.minimum(csp, demand_elec[day])
        used_pvt = np.minimum(pvt, demand_elec[day] - used_csp)
        out[day] = {"CSP": used_csp, "PVT": used_pvt,
                    "unmet": demand_elec[day] - used_csp - used_pvt,
                    "surplus": csp + pvt - used_csp - used_pvt,
                    "surplus_CSP": csp - used_csp, "surplus_PVT": pvt - used_pvt}
    return out


# =============================================================================
# 7. PRINT
# =============================================================================
def print_results(r, data):
    _, _, demand_heat, demand_elec, _ = data
    count_heliostats, count_PVT_panels, count_turbines = r["counts"]
    par = r["par"]
    print("\n=== Assignment 2 design ===")
    print(f"Eff_CSP = {par['eff_CSP']:.4f}   Coef_power = {par['coef_power']:.4f}   "
          f"rotor area = {area_rotor:,.1f} m2")
    print(f"LHV_ethanol = {LHV_ethanol:.3f} MJ/kg  (water {mass_water_per_kg_ethanol:.4f} kg per kg ethanol)")
    print(f"Q6  Count_turbines    = {count_turbines}")
    print(f"Q8  Count_heliostats  = {count_heliostats}   (Assignment 1: {count_heliostats_A1})")
    print(f"Q8  Count_PVT_panels  = {count_PVT_panels:,}   (Assignment 1: {count_PVT_panels_A1:,})")
    area = count_heliostats * area_per_heliostat + count_PVT_panels * area_per_PVT_panel
    print(f"    Collector area    = {area:,.0f} m2 of {area_max:,.0f} m2")
    for day in days:
        u = r["use"]
        demand = demand_elec[day].sum()
        kg = r["ethanol_kg"][day]
        burning = kg > 1e-9
        rankine_out = u["CSP"][day] + u["biomass"][day]
        print(f"\n  {day}: electricity demand {demand:6.1f} GJ, unmet {max(0.0, u['unmet'][day].sum()):.3f} GJ")
        for s in ["CSP", "PVT", "wind", "biomass"]:
            print(f"    {s:<8} delivers {u[s][day].sum():6.1f} GJ ({u[s][day].sum() / demand:5.1%})")
        print(f"    generated {r['avail_CSP'][day].sum() + r['avail_PVT'][day].sum() + r['avail_wind'][day].sum():6.1f} GJ "
              f"(CSP {r['avail_CSP'][day].sum():.1f}, PVT {r['avail_PVT'][day].sum():.1f}, wind {r['avail_wind'][day].sum():.1f}); "
              f"surplus {r['surplus_elec'][day].sum():6.1f} GJ")
        print(f"    heat: PVT {r['supply_heat'][day].sum():6.1f} GJ vs demand sunlit {demand_heat[day][par['sun_hrs'][day]].sum():6.1f} "
              f"/ whole day {demand_heat[day].sum():6.1f} GJ")
        print(f"    Q7 ethanol: {kg.sum():8.0f} kg/day over {int(burning.sum())} h, hours {[int(h) for h in np.where(burning)[0]]}")
        if burning.any():
            print(f"       rate: mean {kg[burning].mean():7.0f} kg/h = {kg[burning].mean() / 3600:.3f} kg/s, "
                  f"peak {kg.max():7.0f} kg/h = {kg.max() / 3600:.3f} kg/s at {int(np.argmax(kg))}:00")
        print(f"    Rankine peak: solar {u['CSP'][day].max():5.1f} GJ/h vs ethanol {u['biomass'][day].max():5.1f} GJ/h "
              f"(combined max {rankine_out.max():5.1f})")
        surplus_by = {"CSP": r["avail_CSP"][day].sum() - u["CSP"][day].sum(),
                      "PVT": r["avail_PVT"][day].sum() - u["PVT"][day].sum(),
                      "wind": r["avail_wind"][day].sum() - u["wind"][day].sum()}
        print("    surplus by source: " + ", ".join(f"{k} {v:.1f}" for k, v in surplus_by.items()) + " GJ")


# =============================================================================
# 8. FIGURES
# =============================================================================
COLOR_CSP = "#2a78d6"      # categorical slot 1
COLOR_PVT = "#eb6834"      # slot 2
COLOR_WIND = "#1baf7a"     # slot 3
COLOR_BIOMASS = "#eda100"  # slot 4
COLOR_DEMAND = "#0b0b0b"
COLOR_SURPLUS = "#b9b8b3"
COLOR_GRID = "#e4e3df"
COLOR_TEXT_2 = "#52514e"
SURFACE = "#fcfcfb"

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10,
    "figure.facecolor": SURFACE, "axes.facecolor": SURFACE,
    "axes.edgecolor": COLOR_TEXT_2, "axes.labelcolor": "#0b0b0b",
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": COLOR_GRID, "grid.linewidth": 0.8,
    "xtick.color": COLOR_TEXT_2, "ytick.color": COLOR_TEXT_2,
    "legend.frameon": False, "hatch.color": "#8d8c87", "hatch.linewidth": 0.6,
})


def tint(color, share_white=0.72):
    """Light version of a colour (blended with the chart surface), used for the hatched surplus segments."""
    from matplotlib.colors import to_rgb
    c, w = np.array(to_rgb(color)), np.array(to_rgb(SURFACE))
    return tuple((1 - share_white) * c + share_white * w)


def tint(color, share_white=0.72):
    """Light version of a colour (blended with the chart surface), used for the hatched surplus segments."""
    from matplotlib.colors import to_rgb
    c, w = np.array(to_rgb(color)), np.array(to_rgb(SURFACE))
    return tuple((1 - share_white) * c + share_white * w)


def figure_dispatch(r, data, path):
    """
    Q9: hourly electricity per technology. For each technology the solid segment is the energy USED to
    meet demand and the hatched segment (same colour, lighter) is the energy PRODUCED BUT NOT USED (surplus).
    The full bar is therefore everything produced in that hour. Heat as in Assignment 1.
    """
    _, _, demand_heat, demand_elec, _ = data
    count_heliostats, count_PVT_panels, count_turbines = r["counts"]
    fig, axes = plt.subplots(2, 2, figsize=(12.5, 8.6), sharex=True)
    hours = np.array(hrs)
    techs = [("CSP", COLOR_CSP, "CSP", "avail_CSP"), ("PVT", COLOR_PVT, "PVT", "avail_PVT"),
             ("wind", COLOR_WIND, "Wind", "avail_wind"), ("biomass", COLOR_BIOMASS, "Ethanol", None)]
    for col, day in enumerate(days):
        ax = axes[0, col]
        bottom = np.zeros(24)
        for s, color, label, _ in techs:                       # used, stacked first
            v = r["use"][s][day]
            ax.bar(hours, v, 0.8, bottom=bottom, color=color, edgecolor=SURFACE, linewidth=1)
            bottom += v
        for s, color, label, key in techs:                     # then the surplus of each technology
            if key is None:
                continue
            v = np.maximum(0.0, r[key][day] - r["use"][s][day])
            ax.bar(hours, v, 0.8, bottom=bottom, facecolor=tint(color), edgecolor=color, hatch="////",
                   linewidth=0.6)
            bottom += v
        ax.step(hours, demand_elec[day], where="mid", color=COLOR_DEMAND, linewidth=2)
        ax.set_title(f"Electricity, {day}", loc="left")
        ax.set_ylabel("Energy per hour (GJ)")

        ax = axes[1, col]
        heat_used = np.minimum(r["supply_heat"][day], demand_heat[day])
        heat_surplus = r["supply_heat"][day] - heat_used
        ax.bar(hours, heat_used, 0.8, color=COLOR_PVT, label="PVT heat: used", edgecolor=SURFACE, linewidth=1)
        ax.bar(hours, heat_surplus, 0.8, bottom=heat_used, facecolor=tint(COLOR_PVT), edgecolor=COLOR_PVT,
               hatch="////", linewidth=0.6, label="PVT heat: surplus")
        ax.step(hours, demand_heat[day], where="mid", color=COLOR_DEMAND, linewidth=2, label="Heat demand")
        ax.set_title(f"Heat, {day} (PVT only)", loc="left")
        ax.set_ylabel("Energy per hour (GJ)")
        ax.set_xlabel("Hour of day")
        ax.set_xticks(range(0, 24, 2))
    top = max(axes[0, 0].get_ylim()[1], axes[0, 1].get_ylim()[1])
    for ax in axes[0]:
        ax.set_ylim(0, top)
    axes[1, 0].legend(loc="upper left", fontsize=8.5)
    from matplotlib.lines import Line2D
    handles = []
    for _, color, label, key in techs:
        handles.append(Patch(color=color, label=f"{label}: used"))
        if key:
            handles.append(Patch(facecolor=tint(color), edgecolor=color, hatch="////", label=f"{label}: surplus"))
    handles.append(Line2D([], [], color=COLOR_DEMAND, linewidth=2, label="Electricity demand"))
    axes[0, 0].legend(handles=handles, loc="upper left", fontsize=8, ncol=2, handlelength=1.6)
    for ax in axes.flat:
        ax.set_axisbelow(True)
    fig.suptitle(f"Hourly plant output vs demand ({count_heliostats} heliostats, {count_PVT_panels:,} PVT panels, "
                 f"{count_turbines} wind turbines, ethanol back-up). Full bar = everything produced in that hour.",
                 x=0.01, ha="left", fontsize=11.5)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def figure_comparison(r, a1, data, path):
    """
    Q9: whole-day electricity, Assignment 1 plant vs Assignment 2 plant, both days.
    Solid = demand met, by source; dashed box = unmet demand; hatched = surplus, by source (same colours).
    Below each bar: share of demand covered and total surplus.
    """
    _, _, _, demand_elec, _ = data
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.2), sharey=True)
    techs = [("CSP", COLOR_CSP), ("PVT", COLOR_PVT), ("wind", COLOR_WIND), ("biomass", COLOR_BIOMASS)]
    for col, day in enumerate(days):
        ax = axes[col]
        demand = demand_elec[day].sum()
        cases = {
            "Assignment 1\n(CSP + PVT)": {
                "used": {"CSP": a1[day]["CSP"].sum(), "PVT": a1[day]["PVT"].sum(), "wind": 0.0, "biomass": 0.0},
                "unmet": a1[day]["unmet"].sum(),
                "surplus": {"CSP": a1[day]["surplus_CSP"].sum(), "PVT": a1[day]["surplus_PVT"].sum(), "wind": 0.0}},
            "Assignment 2\n(+ wind + ethanol)": {
                "used": {s: r["use"][s][day].sum() for s, _ in techs},
                "unmet": max(0.0, r["use"]["unmet"][day].sum()),
                "surplus": {"CSP": r["avail_CSP"][day].sum() - r["use"]["CSP"][day].sum(),
                            "PVT": r["avail_PVT"][day].sum() - r["use"]["PVT"][day].sum(),
                            "wind": r["avail_wind"][day].sum() - r["use"]["wind"][day].sum()}},
        }
        labels = []
        for x, (name, v) in enumerate(cases.items()):
            bottom = 0.0
            for s, color in techs:
                ax.bar(x, v["used"][s], 0.55, bottom=bottom, color=color, edgecolor=SURFACE, linewidth=2)
                bottom += v["used"][s]
            if v["unmet"] > 0.05:
                ax.bar(x, v["unmet"], 0.55, bottom=bottom, facecolor=SURFACE, edgecolor=COLOR_DEMAND,
                       linewidth=1.2, linestyle="--")
                ax.text(x, bottom + v["unmet"] / 2, f"unmet\n{v['unmet']:.0f} GJ", ha="center", va="center",
                        fontsize=8.5)
                bottom += v["unmet"]
            for s, color in techs[:3]:
                if v["surplus"][s] > 0.05:
                    ax.bar(x, v["surplus"][s], 0.55, bottom=bottom, facecolor=tint(color), edgecolor=color,
                           hatch="////", linewidth=0.6)
                    bottom += v["surplus"][s]
            total_surplus = sum(v["surplus"].values())
            covered = sum(v["used"].values())
            ax.text(x, bottom + 30, f"surplus {total_surplus:.0f} GJ", ha="center", fontsize=8.5, color=COLOR_TEXT_2)
            labels.append(f"{name}\ncovered {covered:.0f} GJ ({covered / demand:.0%})\nsurplus {total_surplus:.0f} GJ")
        ax.axhline(demand, color=COLOR_DEMAND, linewidth=1.5)
        ax.text(1.32, demand, f"daily demand\n{demand:.0f} GJ", va="center", fontsize=8.5)
        ax.set_xticks([0, 1], labels)
        ax.set_xlim(-0.5, 1.9)
        ax.set_title(f"Electricity over the whole day, {day}", loc="left")
        ax.grid(axis="x", visible=False)
        ax.set_axisbelow(True)
    axes[0].set_ylabel("Energy per day (GJ)")
    handles = []
    for (s, color), label in zip(techs, ["CSP", "PVT", "Wind", "Ethanol"]):
        handles.append(Patch(color=color, label=f"{label}: demand met"))
        if s != "biomass":
            handles.append(Patch(facecolor=tint(color), edgecolor=color, hatch="////", label=f"{label}: surplus"))
    handles.append(Patch(facecolor=SURFACE, edgecolor=COLOR_DEMAND, linestyle="--", label="Unmet demand"))
    fig.legend(handles=handles, loc="lower center", ncol=8, fontsize=8, handlelength=1.6, columnspacing=1.0)
    fig.suptitle("Demand met vs surplus per day: Assignment 1 plant vs hybrid plant", x=0.01, ha="left", fontsize=12)
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    fig.savefig(path, dpi=200)
    plt.close(fig)


def wind_sweep(r, data):
    """Q10: for 75-95% of the Betz limit, (a) fixed design, dispatched; (b) model re-solved."""
    rows = []
    for frac in np.round(np.arange(0.75, 0.9501, 0.01), 2):
        par = derived(data, frac)
        # (a) same plant, physical output: (C4) and (C5) relaxed, because with the units fixed the
        #     rules can no longer be re-balanced; ethanol simply fills whatever wind does not cover
        counts_fixed, use_fixed = solve(data, par, fixed_counts=r["counts"], wind_rule=False, share_rule=False)
        s_fixed = summarise(counts_fixed, use_fixed, data, par)
        # (b) re-sized plant
        counts_new, use_new = solve(data, par)
        for day in days:
            gen = s_fixed["avail_CSP"][day].sum() + s_fixed["avail_PVT"][day].sum() + s_fixed["avail_wind"][day].sum()
            rows.append({"frac": frac, "day": day,
                         "total_output": gen + use_fixed["biomass"][day].sum(),
                         "wind_generated": s_fixed["avail_wind"][day].sum(),
                         "wind_used": use_fixed["wind"][day].sum(),
                         "ethanol_elec": use_fixed["biomass"][day].sum(),
                         "unmet": use_fixed["unmet"][day].sum(),
                         "CSP_share": s_fixed["avail_CSP"][day].sum() / (gen + use_fixed["biomass"][day].sum()),
                         "turbines_resized": counts_new[2], "heliostats_resized": counts_new[0]})
    return rows


def figure_wind_sweep(sweep, r, data, path):
    """
    Q10: total electricity produced by the same plant (fixed design) at 75-95% of the Betz limit.
    One stacked bar per efficiency step: CSP and PVT (unchanged), wind (grows with Cp), ethanol (what is
    left to fill). The number on top is the total output; the line is the day's electricity demand.
    """
    _, _, _, demand_elec, _ = data
    steps = [0.75, 0.80, 0.85, 0.90, 0.95]
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.0), sharey=True)
    for ax, day in zip(axes, days):
        rows = {round(row["frac"], 2): row for row in sweep if row["day"] == day}
        x = np.arange(len(steps))
        csp = np.full(len(steps), r["avail_CSP"][day].sum())
        pvt = np.full(len(steps), r["avail_PVT"][day].sum())
        wind = np.array([rows[f]["wind_generated"] for f in steps])
        eth = np.array([rows[f]["ethanol_elec"] for f in steps])
        bottom = np.zeros(len(steps))
        for v, color, label in [(csp, COLOR_CSP, "CSP"), (pvt, COLOR_PVT, "PVT"),
                                (wind, COLOR_WIND, "Wind"), (eth, COLOR_BIOMASS, "Ethanol")]:
            ax.bar(x, v, 0.62, bottom=bottom, color=color, edgecolor=SURFACE, linewidth=1, label=label)
            for xi, b, vi in zip(x, bottom, v):
                ax.text(xi, b + vi / 2 if label != "CSP" else b + vi * 0.35, f"{vi:.0f}", ha="center",
                        va="center", fontsize=8.5, color=COLOR_DEMAND if label == "Ethanol" else "white")
            bottom += v
        for xi, t in zip(x, bottom):
            ax.text(xi, t + 20, f"{t:,.0f} GJ", ha="center", va="bottom", fontsize=9.5, fontweight="bold")
        ax.axhline(demand_elec[day].sum(), color=COLOR_DEMAND, linewidth=1.6, linestyle="--")
        ax.text(-0.42, demand_elec[day].sum(), f" demand {demand_elec[day].sum():.0f} GJ ",
                ha="left", va="center", fontsize=8.5, color=COLOR_DEMAND,
                bbox=dict(boxstyle="round,pad=0.25", facecolor=SURFACE, edgecolor=COLOR_DEMAND, linewidth=0.8))
        ax.set_xticks(x)
        ax.set_xticklabels([f"{f * 100:.0f}%" + ("\n(design)" if abs(f - frac_of_betz) < 1e-9 else "")
                            for f in steps])
        ax.set_xlabel("Turbine efficiency (% of the Betz limit)")
        ax.set_title(f"{day}", loc="left")
        ax.set_axisbelow(True)
        ax.grid(axis="x", visible=False)
        ax.set_ylim(0, bottom.max() * 1.12)
    axes[0].set_ylabel("Electricity produced per day (GJ)")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=4, fontsize=9.5, bbox_to_anchor=(0.5, 0.0))
    fig.suptitle(f"Total electricity produced per day by the same plant ({r['counts'][2]} turbines) "
                 f"at different turbine efficiencies", x=0.01, ha="left", fontsize=11.5)
    fig.tight_layout(rect=(0, 0.06, 1, 0.96))
    fig.savefig(path, dpi=200)
    plt.close(fig)


def wind_tradeoff(data, par, counts):
    """
    Trade-off behind the wind rule: for a fixed number of turbines (0-80), the rest of the plant is
    optimised without (C5). Shows how wind replaces ethanol and how much surplus that costs.
    """
    rows = []
    for turbines in sorted(set(range(0, 81, 2)) | {counts[2]}):
        c, use = solve(data, par, fixed_turbines=turbines, wind_rule=False)
        s = summarise(c, use, data, par)
        for day in days:
            rows.append({"turbines": turbines, "day": day, "heliostats": c[0],
                         "surplus": s["surplus_elec"][day].sum(),
                         "wind_used": use["wind"][day].sum(), "ethanol_elec": use["biomass"][day].sum()})
    return rows


def figure_wind_vs_ethanol(sweep, r, path):
    """
    Q10, complements the stacked bars: for the same plant, electricity produced by wind and by
    ethanol as turbine efficiency rises from 75% to 95% of the Betz limit.
    """
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 4.8))
    for ax, day in zip(axes, days):
        d = sorted([row for row in sweep if row["day"] == day], key=lambda row: row["frac"])
        x = np.array([row["frac"] * 100 for row in d])
        wind = np.array([row["wind_generated"] for row in d])
        eth = np.array([row["ethanol_elec"] for row in d])
        ax.plot(x, wind, color=COLOR_WIND, linewidth=2.4, marker="o", markersize=4, label="Wind")
        ax.plot(x, eth, color=COLOR_BIOMASS, linewidth=2.4, marker="o", markersize=4, label="Ethanol")
        for arr, color in [(wind, COLOR_WIND), (eth, COLOR_BIOMASS)]:
            ax.text(x[0] - 0.4, arr[0], f"{arr[0]:.0f}", ha="right", va="center", fontsize=9, color=color,
                    fontweight="bold")
            ax.text(x[-1] + 0.4, arr[-1], f"{arr[-1]:.0f}", ha="left", va="center", fontsize=9, color=color,
                    fontweight="bold")
        ax.axvline(frac_of_betz * 100, color=COLOR_TEXT_2, linewidth=0.8, linestyle=":")
        diff = wind - eth
        cross = np.where((diff[:-1] < 0) & (diff[1:] >= 0))[0]
        if len(cross):
            i = cross[0]
            xc = x[i] + (x[i + 1] - x[i]) * (-diff[i]) / (diff[i + 1] - diff[i])
            yc = np.interp(xc, x, wind)
            ax.plot(xc, yc, marker="o", markersize=9, markerfacecolor=SURFACE, markeredgecolor=COLOR_DEMAND,
                    markeredgewidth=1.5, zorder=5)
            ax.annotate(f"wind overtakes ethanol\nat ≈{xc:.0f}% of Betz", (xc, yc), xytext=(45, -45), textcoords="offset points",
                        fontsize=8.5, ha="center", arrowprops=dict(arrowstyle="-", color=COLOR_TEXT_2, lw=0.8))
        else:
            ax.text(x.mean(), (wind.min() + eth.max()) / 2, "wind above ethanol at every efficiency",
                    ha="center", va="center", fontsize=8.5, color=COLOR_TEXT_2)
        ax.set_xticks([75, 80, 85, 90, 95])
        ax.set_xticklabels(["75%", "80%", "85%\n(design)", "90%", "95%"])
        ax.set_xlim(72, 98)
        ax.set_xlabel("Turbine efficiency (% of the Betz limit)")
        ax.set_title(day, loc="left")
        ax.set_axisbelow(True)
        lo, hi = min(wind.min(), eth.min()), max(wind.max(), eth.max())
        pad = max(6.0, 0.12 * (hi - lo))
        ax.set_ylim(np.floor((lo - pad) / 5) * 5, np.ceil((hi + pad) / 5) * 5)
    axes[0].set_ylabel("Electricity produced per day (GJ)")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=2,
               fontsize=9.5, bbox_to_anchor=(0.5, 0.0))
    fig.suptitle(f"Wind vs ethanol electricity produced by the same plant ({r['counts'][2]} turbines) at different "
                 f"turbine efficiencies", x=0.01, ha="left", fontsize=11.5)
    fig.tight_layout(rect=(0, 0.07, 1, 0.96))
    fig.savefig(path, dpi=200)
    plt.close(fig)


def export_workbook(r, a1, sweep, trade, data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    irradiance, T_ambient, demand_heat, demand_elec, speed_wind = data
    F = "Arial"
    f_title, f_sub = Font(name=F, bold=True, size=13), Font(name=F, italic=True, color="555555")
    f_head, f_bold, f_text = Font(name=F, bold=True, color="FFFFFF"), Font(name=F, bold=True), Font(name=F)
    f_input = Font(name=F, color="0000FF")
    fill_head, fill_total = PatternFill("solid", fgColor="2E3A48"), PatternFill("solid", fgColor="EEF1F5")
    fill_design = PatternFill("solid", fgColor="FFFF00")

    def header(ws, row, labels, height=48):
        for c, label in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=c, value=label)
            cell.font, cell.fill = f_head, fill_head
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[row].height = height

    wb = Workbook()
    ws = wb.active
    ws.title = "Design"
    ws["A1"], ws["A1"].font = "Assignment 2: hybrid plant design and key results", f_title
    ws["A2"], ws["A2"].font = "Blue = given or assumed input. Yellow = result of the optimisation.", f_sub
    header(ws, 4, ["Quantity", "Value", "Unit", "Symbol in the mathematical model"], 24)
    count_heliostats, count_PVT_panels, count_turbines = r["counts"]
    items = [
        ("Inputs", None, None, None, None),
        ("Turbine efficiency as fraction of the Betz limit", frac_of_betz, "fraction", "Frac_of Betz", "in"),
        ("Rotor radius", radius_rotor, "m", "Radius_rotor", "in"),
        ("Air density (assumed, sea level)", density_air, "kg/m³", "Density_air", "in"),
        ("Ethanol higher heating value", HHV_ethanol, "MJ/kg", "HHV_ethanol", "in"),
        ("Heat of evaporation of water", H_evap_water, "MJ/kg", "H_evap water", "in"),
        ("Water formed per kg of ethanol (from Q2)", mass_water_per_kg_ethanol, "kg/kg", "Mass_water per kg ethanol", "in"),
        ("Ethanol lower heating value", LHV_ethanol, "MJ/kg", "LHV_ethanol", "in"),
        ("Rankine cycle efficiency", eff_rankine, "fraction", "Eff_rankine", "in"),
        ("Minimum CSP share of all electricity produced (surplus included)", share_CSP_min, "fraction", "Share_CSP,min", "in"),
        ("Results", None, None, None, None),
        ("Number of wind turbines (Q6)", count_turbines, "units", "Count_turbines", "out"),
        ("Number of heliostats (Q8)", count_heliostats, "units", "Count_heliostats", "out"),
        ("Number of PVT panels (Q8)", count_PVT_panels, "units", "Count_PVT panels", "out"),
        ("Heliostats in Assignment 1", count_heliostats_A1, "units", "", "in"),
        ("PVT panels in Assignment 1", count_PVT_panels_A1, "units", "", "in"),
    ]
    row = 5
    for label, value, unit, symbol, kind in items:
        ws.cell(row=row, column=1, value=label).font = f_bold if value is None else f_text
        if value is not None:
            v = ws.cell(row=row, column=2, value=value)
            v.font = f_input
            if kind == "out":
                v.fill = fill_design
            ws.cell(row=row, column=3, value=unit).font = f_text
            ws.cell(row=row, column=4, value=symbol).font = f_sub
        row += 1
    for col, w in zip("ABCD", (52, 14, 12, 30)):
        ws.column_dimensions[col].width = w

    cols = ["Hour", "Irradiance\n[W/m²]", "Wind speed\n[m/s]", "Electricity demand\n[GJ]",
            "CSP generated\n[GJ]", "PVT generated\n[GJ]", "Wind generated\n[GJ]",
            "CSP used\n[GJ]", "PVT used\n[GJ]", "Wind used\n[GJ]", "Ethanol-fired Rankine\n[GJ]",
            "Unmet demand\n[GJ]", "Surplus\n[GJ]", "Ethanol burned\n[kg/h]", "Ethanol burned\n[kg/s]",
            "PVT heat\n[GJ]", "Heat demand\n[GJ]", "Assignment 1: unmet\n[GJ]"]
    for day in days:
        ws = wb.create_sheet(f"{day} hourly")
        ws["A1"], ws["A1"].font = f"{day} 2019: hourly dispatch of the hybrid plant", f_title
        ws["A2"], ws["A2"].font = ("Generated = what the installed units can produce. Used = delivered to demand. "
                                   "Surplus = generated − used (column formula)."), f_sub
        header(ws, 3, cols, 62)
        u = r["use"]
        for hr in hrs:
            rr = 4 + hr
            vals = [hr, irradiance[day][hr], speed_wind[day][hr], demand_elec[day][hr],
                    r["avail_CSP"][day][hr], r["avail_PVT"][day][hr], r["avail_wind"][day][hr],
                    u["CSP"][day][hr], u["PVT"][day][hr], u["wind"][day][hr], u["biomass"][day][hr],
                    u["unmet"][day][hr], f"=E{rr}+F{rr}+G{rr}-H{rr}-I{rr}-J{rr}",
                    r["ethanol_kg"][day][hr], f"=N{rr}/3600",
                    r["supply_heat"][day][hr], demand_heat[day][hr], a1[day]["unmet"][hr]]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=rr, column=c, value=(round(float(v), 6) if not isinstance(v, str) else v))
                cell.font = f_text
                cell.number_format = "0" if c == 1 else ("0.000" if c == 15 else ("0" if c == 14 else "0.00"))
        tr = 28
        ws.cell(row=tr, column=1, value="Day total").font = f_bold
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=tr, column=c)
            cell.fill = fill_total
            if c >= 4 and c not in (15,):
                L = get_column_letter(c)
                cell.value, cell.font = f"=SUM({L}4:{L}27)", f_bold
                cell.number_format = "#,##0" if c == 14 else "0.0"
        ws.column_dimensions["A"].width = 10
        for c in range(2, len(cols) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.freeze_panes = "B4"

    ws = wb.create_sheet("Q10 wind efficiency")
    ws["A1"], ws["A1"].font = "Q10: turbine efficiency between 75% and 95% of the Betz limit", f_title
    ws["A2"], ws["A2"].font = ("Columns C-G: same plant as designed, dispatched without re-balancing the rules. "
                               "Columns H-I: plant re-solved at that efficiency."), f_sub
    header(ws, 3, ["Efficiency\n[% of Betz]", "Day", "Total electricity output\n[GJ/day]",
                   "Wind generated\n[GJ/day]", "Wind used\n[GJ/day]", "Ethanol-fired electricity\n[GJ/day]",
                   "CSP share of all\nelectricity produced", "Turbines needed\nif re-sized",
                   "Heliostats needed\nif re-sized"], 48)
    for s in sweep:
        ws.append([round(s["frac"] * 100), s["day"], round(s["total_output"], 2), round(s["wind_generated"], 2),
                   round(s["wind_used"], 2), round(s["ethanol_elec"], 2), round(s["CSP_share"], 4),
                   s["turbines_resized"], s["heliostats_resized"]])
        ws.cell(ws.max_row, 7).number_format = "0.0%"
    for c, w in zip("ABCDEFGHI", (14, 12, 18, 16, 14, 18, 18, 16, 16)):
        ws.column_dimensions[c].width = w

    ws = wb.create_sheet("Wind-ethanol trade-off")
    ws["A1"], ws["A1"].font = "Wind vs ethanol: what each extra turbine changes", f_title
    ws["A2"], ws["A2"].font = ("Turbine count fixed; heliostats, PVT and dispatch optimised without the wind >= "
                               "ethanol rule. The chosen design is where wind >= ethanol on both days."), f_sub
    header(ws, 3, ["Wind turbines", "Day", "Heliostats", "Wind delivered\n[GJ/day]",
                   "Ethanol-fired electricity\n[GJ/day]", "Surplus\n[GJ/day]"], 48)
    for t in trade:
        ws.append([t["turbines"], t["day"], t["heliostats"], round(t["wind_used"], 2),
                   round(t["ethanol_elec"], 2), round(t["surplus"], 2)])
        if t["turbines"] == r["counts"][2]:
            for c in range(1, 7):
                ws.cell(row=ws.max_row, column=c).fill = fill_design
    for c, w in zip("ABCDEF", (14, 12, 12, 16, 20, 14)):
        ws.column_dimensions[c].width = w
    wb.save(path)


# =============================================================================
# MAIN
# =============================================================================
if __name__ == "__main__":
    project_dir = Path(__file__).resolve().parent.parent
    folder_figures, folder_data = project_dir / "4_Visualizations", project_dir / "5_Data"
    folder_figures.mkdir(exist_ok=True)
    data = read_data(folder_data / "Input_Data_Assignment_1.xlsx", folder_data / "Input_Data_Assignment_2_Wind.xlsx")

    par = derived(data)
    counts, use = solve(data, par)
    result = summarise(counts, use, data, par)
    print_results(result, data)

    a1 = assignment1_balance(data, par)
    for day in days:
        print(f"  Assignment 1 plant, {day}: unmet {a1[day]['unmet'].sum():.1f} GJ, surplus {a1[day]['surplus'].sum():.1f} GJ")

    sweep = wind_sweep(result, data)
    print("\nQ10 sweep (June / December total output, GJ; turbines if re-sized):")
    for frac in sorted({s["frac"] for s in sweep}):
        j, d = [s for s in sweep if s["frac"] == frac]
        print(f"  {frac:.2f}: {j['total_output']:7.1f} / {d['total_output']:7.1f}   "
              f"wind {j['wind_generated']:6.1f} / {d['wind_generated']:6.1f}   "
              f"ethanol {j['ethanol_elec']:5.1f} / {d['ethanol_elec']:5.1f}   turbines {j['turbines_resized']}")

    figure_dispatch(result, data, folder_figures / "Q9_Hourly_Supply_vs_Demand.png")
    figure_comparison(result, a1, data, folder_figures / "Q9_Comparison_with_Assignment_1.png")
    figure_wind_sweep(sweep, result, data, folder_figures / "Q10_Wind_Efficiency_Sweep.png")

    trade = wind_tradeoff(data, par, result["counts"])
    print("\nWind-ethanol trade-off (June | December): turbines, surplus, wind, ethanol")
    for t in sorted({row["turbines"] for row in trade}):
        j, d = [row for row in trade if row["turbines"] == t]
        if t % 10 == 0 or t == result["counts"][2]:
            print(f"  {t:3d}: {j['surplus']:6.0f} {j['wind_used']:5.1f} {j['ethanol_elec']:5.1f} | "
                  f"{d['surplus']:6.0f} {d['wind_used']:5.1f} {d['ethanol_elec']:5.1f}")
    figure_wind_vs_ethanol(sweep, result, folder_figures / "Q10_Wind_vs_Ethanol.png")
    export_workbook(result, a1, sweep, trade, data, folder_data / "Results_Hybrid_Plant_A2.xlsx")
    print(f"\nSaved figures to {folder_figures} and results to {folder_data}")
